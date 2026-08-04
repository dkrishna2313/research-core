"""Tests for content extractors."""

from __future__ import annotations

import io
import sys
from datetime import UTC, datetime
from unittest.mock import MagicMock, patch

import pytest

pytestmark = pytest.mark.web

from research_core.adapters.web.extract import (  # noqa: E402
    CompositeExtractor,
    DocxExtractor,
    PlainTextExtractor,
    PyPDFExtractor,
    TrafilaturaExtractor,
    default_extractor,
)
from research_core.adapters.web.models import FetchedResource  # noqa: E402
from research_core.exceptions import ProviderExecutionError  # noqa: E402

FIXED_TS = datetime(2024, 3, 1, tzinfo=UTC)


def _resource(
    content: bytes = b"",
    content_type: str = "text/html",
    url: str = "https://example.com",
) -> FetchedResource:
    return FetchedResource(
        requested_url=url,
        final_url=url,
        status_code=200,
        content_type=content_type,
        content=content,
        retrieved_at=FIXED_TS,
    )


class TestTrafilaturaExtractor:
    def test_can_extract_html(self) -> None:
        ext = TrafilaturaExtractor()
        assert ext.can_extract("text/html; charset=utf-8", "https://example.com") is True

    def test_can_extract_xhtml(self) -> None:
        ext = TrafilaturaExtractor()
        assert ext.can_extract("application/xhtml+xml", "https://example.com") is True

    def test_cannot_extract_pdf(self) -> None:
        ext = TrafilaturaExtractor()
        assert ext.can_extract("application/pdf", "https://example.com/doc.pdf") is False

    def test_extracts_html_text(self) -> None:
        ext = TrafilaturaExtractor()
        html = b"<html><head><title>My Title</title></head><body><p>Hello world.</p></body></html>"
        resource = _resource(content=html)

        mock_trafilatura = MagicMock()
        mock_trafilatura.extract.return_value = "Hello world."
        with patch.dict(sys.modules, {"trafilatura": mock_trafilatura}):
            result = ext.extract(resource)

        assert "Hello world." in result.text
        assert result.extraction_method == "trafilatura"

    def test_title_extracted_from_html_tag(self) -> None:
        ext = TrafilaturaExtractor()
        html = (
            b"<html><head><title>Page Title</title></head>"
            b"<body><p>Content here.</p></body></html>"
        )
        resource = _resource(content=html)

        mock_trafilatura = MagicMock()
        mock_trafilatura.extract.return_value = "Content here."
        with patch.dict(sys.modules, {"trafilatura": mock_trafilatura}):
            result = ext.extract(resource)

        assert result.title == "Page Title"

    def test_empty_result_raises(self) -> None:
        ext = TrafilaturaExtractor()
        html = b"<nav>Menu only</nav>"
        resource = _resource(content=html)

        mock_trafilatura = MagicMock()
        mock_trafilatura.extract.return_value = ""
        with (
            patch.dict(sys.modules, {"trafilatura": mock_trafilatura}),
            pytest.raises(ProviderExecutionError, match="empty text"),
        ):
            ext.extract(resource)

    def test_missing_trafilatura_raises(self) -> None:
        ext = TrafilaturaExtractor()
        resource = _resource(content=b"<html>test</html>")
        with (
            patch.dict(sys.modules, {"trafilatura": None}),  # type: ignore[dict-item]
            pytest.raises(ProviderExecutionError, match="trafilatura"),
        ):
            ext.extract(resource)


class TestPyPDFExtractor:
    def test_can_extract_pdf_content_type(self) -> None:
        ext = PyPDFExtractor()
        assert ext.can_extract("application/pdf", "https://example.com") is True

    def test_can_extract_pdf_url(self) -> None:
        ext = PyPDFExtractor()
        assert ext.can_extract("application/octet-stream", "https://example.com/doc.pdf") is True

    def test_cannot_extract_html(self) -> None:
        ext = PyPDFExtractor()
        assert ext.can_extract("text/html", "https://example.com/page.html") is False

    def test_extracts_pdf_pages(self) -> None:
        # Build a minimal real PDF in memory using pypdf if available
        try:
            from pypdf import PdfWriter

            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            buf = io.BytesIO()
            writer.write(buf)
            pdf_bytes = buf.getvalue()
        except ImportError:
            pytest.skip("pypdf not installed")

        ext = PyPDFExtractor()
        resource = _resource(content=pdf_bytes, content_type="application/pdf")
        # Blank page will have no text — so test parse success, not text content
        try:
            result = ext.extract(resource)
            assert result.extraction_method == "pypdf"
            assert result.page_count == 1
        except ProviderExecutionError as exc:
            assert "empty text" in str(exc)

    def test_missing_pypdf_raises(self) -> None:
        ext = PyPDFExtractor()
        resource = _resource(content=b"%PDF-test", content_type="application/pdf")
        with (
            patch.dict(sys.modules, {"pypdf": None}),  # type: ignore[dict-item]
            pytest.raises(ProviderExecutionError, match="pypdf"),
        ):
            ext.extract(resource)

    def test_malformed_pdf_raises(self) -> None:
        try:
            from pypdf import PdfReader  # noqa: F401
        except ImportError:
            pytest.skip("pypdf not installed")

        ext = PyPDFExtractor()
        resource = _resource(content=b"not a pdf at all", content_type="application/pdf")
        with pytest.raises(ProviderExecutionError):
            ext.extract(resource)


class TestDocxExtractor:
    def test_can_extract_docx_content_type(self) -> None:
        ext = DocxExtractor()
        assert (
            ext.can_extract(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "https://example.com",
            )
            is True
        )

    def test_can_extract_docx_url(self) -> None:
        ext = DocxExtractor()
        assert ext.can_extract("application/octet-stream", "https://example.com/doc.docx") is True

    def test_cannot_extract_html(self) -> None:
        ext = DocxExtractor()
        assert ext.can_extract("text/html", "https://example.com") is False

    def test_missing_docx_raises(self) -> None:
        ext = DocxExtractor()
        resource = _resource(content=b"PK...", content_type="application/vnd.openxmlformats")
        with (
            patch.dict(sys.modules, {"docx": None}),  # type: ignore[dict-item]
            pytest.raises(ProviderExecutionError, match="python-docx"),
        ):
            ext.extract(resource)

    def test_extracts_docx_paragraphs(self) -> None:
        try:
            import docx as docx_module

            doc = docx_module.Document()
            doc.add_paragraph("Hello from DOCX.")
            doc.add_paragraph("Second paragraph.")
            buf = io.BytesIO()
            doc.save(buf)
            docx_bytes = buf.getvalue()
        except ImportError:
            pytest.skip("python-docx not installed")

        ext = DocxExtractor()
        resource = _resource(
            content=docx_bytes,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        result = ext.extract(resource)
        assert "Hello from DOCX." in result.text
        assert "Second paragraph." in result.text
        assert result.extraction_method == "docx"


class TestPlainTextExtractor:
    def test_can_extract_text_plain(self) -> None:
        ext = PlainTextExtractor()
        assert ext.can_extract("text/plain; charset=utf-8", "https://example.com") is True

    def test_cannot_extract_html(self) -> None:
        ext = PlainTextExtractor()
        assert ext.can_extract("text/html", "https://example.com") is False

    def test_extracts_utf8_text(self) -> None:
        ext = PlainTextExtractor()
        content = b"Hello, world!"
        resource = _resource(content=content, content_type="text/plain")
        result = ext.extract(resource)
        assert result.text == "Hello, world!"
        assert result.extraction_method == "plaintext"

    def test_empty_text_raises(self) -> None:
        ext = PlainTextExtractor()
        resource = _resource(content=b"   ", content_type="text/plain")
        with pytest.raises(ProviderExecutionError, match="empty"):
            ext.extract(resource)

    def test_invalid_bytes_replaced(self) -> None:
        ext = PlainTextExtractor()
        # Invalid UTF-8 bytes
        resource = _resource(content=b"Hello \xff\xfe world", content_type="text/plain")
        result = ext.extract(resource)
        assert "Hello" in result.text
        assert "world" in result.text


class TestCompositeExtractor:
    def test_selects_pdf_extractor_for_pdf(self) -> None:
        ext = default_extractor()
        assert ext.can_extract("application/pdf", "https://example.com/doc.pdf") is True

    def test_selects_html_extractor_for_html(self) -> None:
        ext = default_extractor()
        assert ext.can_extract("text/html", "https://example.com") is True

    def test_selects_plaintext_extractor(self) -> None:
        ext = default_extractor()
        assert ext.can_extract("text/plain", "https://example.com") is True

    def test_unsupported_content_type_raises(self) -> None:
        ext = CompositeExtractor([PlainTextExtractor()])
        resource = _resource(content=b"binary", content_type="application/octet-stream")
        with pytest.raises(ProviderExecutionError, match="unsupported content type"):
            ext.extract(resource)

    def test_default_extractor_construction_does_not_import_deps(self) -> None:
        with patch.dict(
            sys.modules,
            {
                "pypdf": None,  # type: ignore[dict-item]
                "docx": None,  # type: ignore[dict-item]
                "trafilatura": None,  # type: ignore[dict-item]
            },
        ):
            ext = default_extractor()
        assert ext is not None
